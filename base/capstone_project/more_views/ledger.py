# Updated ledger.py with public ledger view and restricted downloads
from capstone_project.models import blockchain, Donation
from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from django.views.decorators.cache import never_cache
from django.core.paginator import Paginator
from django.contrib import messages
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
import io
from io import BytesIO
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch
from django.http import FileResponse
import logging
from dateutil import parser as date_parser

logger = logging.getLogger(__name__)

def generate_receipt_pdf(donation):
    buffer = BytesIO()
    p = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter

    p.setFont("Helvetica-Bold", 16)
    p.drawString(inch, height - inch, "Knights of Columbus - Donation Receipt")

    # Use unmasked/full real information for private receipts
    donor_name = f"{donation.first_name or ''} {donation.middle_initial or ''} {donation.last_name or ''}".strip() or "Anonymous Donor"
    email = donation.email or "N/A"
    event_name = donation.event.name if donation.event else "General Donation"

    data = {
        'transaction_id': donation.transaction_id,
        'donor_name': donor_name,
        'email': email,
        'amount': donation.amount,
        'donation_date': donation.donation_date,
        'payment_method': donation.payment_method.capitalize(),
        'event_name': event_name,
        'status': donation.get_status_display(),
        'block_index': "Pending" if donation.status != 'completed' else "Recorded",  # Simple for PDF
    }

    y = height - 2 * inch
    p.setFont("Helvetica", 12)
    for key, value in data.items():
        p.drawString(inch, y, f"{key.replace('_', ' ').title()}: {value}")
        y -= 0.25 * inch

    p.drawString(inch, inch, f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    p.drawString(inch, inch - 0.25 * inch, "Thank you for your support! Verify on blockchain ledger.")

    p.save()
    buffer.seek(0)
    return buffer

@login_required
def download_receipt(request, donation_id):
    donation = get_object_or_404(Donation, id=donation_id)
    
    # Authorization check: Admins/officers can download any; users can download their own if email matches
    is_admin_or_officer = request.user.role in ['admin', 'officer']
    is_own_donation = request.user.is_authenticated and request.user.email == donation.email
    
    if not (is_admin_or_officer or is_own_donation):
        messages.info(request, "You can request the receipt via email if this is your donation.")
        return redirect('request_receipt', donation_id=donation_id)
    
    # Optional council restriction for officers (keep if desired)
    if request.user.role == 'officer' and donation.council and donation.council != request.user.council:
        messages.error(request, "You can only download receipts for your council's donations.")
        return redirect('blockchain')
    
    # Log the download
    logger.info(f"User {request.user.username} (role: {request.user.role}, council: {request.user.council.name if request.user.council else 'None'}) downloaded receipt for donation {donation_id} (council: {donation.council.name if donation.council else 'None'})")
    
    pdf_buffer = generate_receipt_pdf(donation)
    return FileResponse(
        pdf_buffer,
        as_attachment=True,
        filename=f"receipt_{donation.transaction_id}.pdf"
    )

def mask_name(name):
    """Partially mask a name for privacy: reveal first/last letter per word (>2 chars), star middle."""
    if not name:
        return 'N/A'
    parts = name.split()
    masked_parts = []
    for part in parts:
        if len(part) <= 2:
            masked_parts.append(part)  # Keep short parts/initials as-is
        else:
            middle_len = len(part) - 2
            masked_parts.append(part[0] + '*' * middle_len + part[-1])
    return ' '.join(masked_parts)

def mask_email(email):
    """Partially mask email local part: reveal first/last, star middle; keep domain."""
    if not email or '@' not in email:
        return 'N/A'
    local, domain = email.split('@', 1)
    if len(local) <= 2:
        masked_local = local
    else:
        middle_len = len(local) - 2
        masked_local = local[0] + '*' * middle_len + local[-1]
    return masked_local + '@' + domain

def normalize_blockchain_data(full_chain, pending_transactions):
    for block in full_chain:
        if not isinstance(block, dict):
            block = {
                'index': block.index,
                'timestamp': block.timestamp,
                'transactions': block.transactions if hasattr(block, 'transactions') else [],
                'proof': block.proof,
                'previous_hash': block.previous_hash,
                'hash': block.hash
            }
        if isinstance(block.get('timestamp'), str):
            try:
                block['timestamp'] = date_parser.parse(block['timestamp'])
            except ValueError:
                block['timestamp'] = None
        for tx in block.get('transactions', []):
            if not isinstance(tx, dict):
                tx_id = getattr(tx, 'id', None)
                is_anonymous = getattr(tx, 'is_anonymous', False)
                donor = tx.get_display_name() if hasattr(tx, 'get_display_name') else getattr(tx, 'donor', 'Anonymous')
                email = getattr(tx, 'email', 'N/A')
                tx = {
                    'id': tx_id,
                    'transaction_id': getattr(tx, 'transaction_id', ''),
                    'donor': donor,
                    'email': email,
                    'amount': getattr(tx, 'amount', '0.00'),
                    'donation_date': getattr(tx, 'donation_date', None),
                    'payment_method': getattr(tx, 'payment_method', 'N/A'),
                    'is_anonymous': is_anonymous,
                    'submitted_by': getattr(tx, 'submitted_by', 'N/A'),
                    'reviewed_by': getattr(tx, 'reviewed_by', 'N/A')
                }
            else:
                tx['id'] = tx.get('id', None)
                if tx['id'] is None and 'transaction_id' in tx:
                    try:
                        donation = Donation.objects.get(transaction_id=tx['transaction_id'])
                        tx['id'] = donation.id
                    except (Donation.DoesNotExist, Donation.MultipleObjectsReturned):
                        tx['id'] = None
                donor_raw = tx.get('donor_name', tx.get('donor', ''))
                is_anonymous = tx.get('is_anonymous', donor_raw == 'Anonymous Donor')
            if is_anonymous:
                tx['donor'] = "Anonymous Donor"
                tx['email'] = "N/A"
            else:
                tx['donor'] = mask_name(donor_raw)
                tx['email'] = mask_email(tx.get('email', ''))
            if 'date' in tx and 'donation_date' not in tx:
                tx['donation_date'] = tx['date']
            if tx.get('donation_date') and isinstance(tx.get('donation_date'), str):
                try:
                    tx['donation_date'] = datetime.strptime(tx['donation_date'], '%Y-%m-%d').date()
                except ValueError:
                    tx['donation_date'] = None
            if 'timestamp' in tx and isinstance(tx['timestamp'], str):
                try:
                    tx['timestamp'] = date_parser.parse(tx['timestamp'])
                except ValueError:
                    tx['timestamp'] = None
            if 'amount' in tx:
                if isinstance(tx['amount'], str):
                    try:
                        tx['amount'] = float(tx['amount'].replace('₱', '').replace(',', ''))
                    except ValueError:
                        tx['amount'] = 0.0
                elif tx['amount'] is None:
                    tx['amount'] = 0.0
            tx['submitted_by'] = tx.get('submitted_by', 'N/A')
            tx['reviewed_by'] = tx.get('reviewed_by', 'N/A')

    for tx in pending_transactions:
        if not isinstance(tx, dict):
            tx_id = getattr(tx, 'id', None)
            is_anonymous = getattr(tx, 'is_anonymous', False)
            donor = tx.get_display_name() if hasattr(tx, 'get_display_name') else getattr(tx, 'donor', 'Anonymous')
            email = getattr(tx, 'email', 'N/A')
            tx = {
                'id': tx_id,
                'transaction_id': getattr(tx, 'transaction_id', ''),
                'donor': donor,
                'email': email,
                'amount': getattr(tx, 'amount', '0.00'),
                'donation_date': getattr(tx, 'donation_date', None),
                'payment_method': getattr(tx, 'payment_method', 'N/A'),
                'is_anonymous': is_anonymous,
                'submitted_by': getattr(tx, 'submitted_by', 'N/A'),
                'reviewed_by': getattr(tx, 'reviewed_by', 'N/A')
            }
        else:
            tx['id'] = tx.get('id', None)
            if tx['id'] is None and 'transaction_id' in tx:
                try:
                    donation = Donation.objects.get(transaction_id=tx['transaction_id'])
                    tx['id'] = donation.id
                except (Donation.DoesNotExist, Donation.MultipleObjectsReturned):
                    tx['id'] = None
                donor_raw = tx.get('donor_name', tx.get('donor', ''))
                is_anonymous = tx.get('is_anonymous', donor_raw == 'Anonymous Donor')
            if is_anonymous:
                tx['donor'] = "Anonymous Donor"
                tx['email'] = "N/A"
            else:
                tx['donor'] = mask_name(donor_raw)
                tx['email'] = mask_email(tx.get('email', ''))
            if 'date' in tx and 'donation_date' not in tx:
                tx['donation_date'] = tx['date']
            if tx.get('donation_date') and isinstance(tx.get('donation_date'), str):
                try:
                    tx['donation_date'] = datetime.strptime(tx['donation_date'], '%Y-%m-%d').date()
                except ValueError:
                    tx['donation_date'] = None
            if 'timestamp' in tx and isinstance(tx['timestamp'], str):
                try:
                    tx['timestamp'] = date_parser.parse(tx['timestamp'])
                except ValueError:
                    tx['timestamp'] = None
            if 'amount' in tx:
                if isinstance(tx['amount'], str):
                    try:
                        tx['amount'] = float(tx['amount'].replace('₱', '').replace(',', ''))
                    except ValueError:
                        tx['amount'] = 0.0
                elif tx['amount'] is None:
                    tx['amount'] = 0.0
            tx['submitted_by'] = tx.get('submitted_by', 'N/A')
            tx['reviewed_by'] = tx.get('reviewed_by', 'N/A')

    return full_chain, pending_transactions

def get_matches_filter(search, date_from, date_to, amount_min, amount_max, method):
    def matches_filter(tx):
        match = True
        if search and not (search.lower() in tx.get('donor', '').lower() or search.lower() in tx.get('transaction_id', '').lower()):
            match = False
        if date_from and tx.get('donation_date') and tx['donation_date'] < datetime.strptime(date_from, '%Y-%m-%d').date():
            match = False
        if date_to and tx.get('donation_date') and tx['donation_date'] > datetime.strptime(date_to, '%Y-%m-%d').date():
            match = False
        if amount_min and amount_min.strip():
            try:
                if tx.get('amount', 0) < float(amount_min):
                    match = False
            except ValueError:
                pass
        if amount_max and amount_max.strip():
            try:
                if tx.get('amount', 0) > float(amount_max):
                    match = False
            except ValueError:
                pass
        if method and method != tx.get('payment_method'):
            match = False
        return match
    return matches_filter

@never_cache
def get_blockchain_data(request):  # Removed @login_required for public transparency
    try:
        full_chain = blockchain.get_chain()
        if not blockchain.is_chain_valid():
            messages.error(request, "Blockchain data is corrupted. Contact support.")
            return redirect('donations')
        
        pending_transactions = blockchain.pending_transactions

        full_chain, pending_transactions = normalize_blockchain_data(full_chain, pending_transactions)

        search = request.GET.get('search', '').lower()
        date_from = request.GET.get('date_from')
        date_to = request.GET.get('date_to')
        amount_min = request.GET.get('amount_min', '')
        amount_max = request.GET.get('amount_max', '')
        method = request.GET.get('method')

        matches_filter_func = get_matches_filter(search, date_from, date_to, amount_min, amount_max, method)

        filtered_chain = []
        for block in full_chain:
            filtered_txs = [tx for tx in block.get('transactions', []) if matches_filter_func(tx)]
            if filtered_txs:
                block_copy = block.copy()
                block_copy['transactions'] = filtered_txs
                filtered_chain.append(block_copy)

        filtered_pending = [tx for tx in pending_transactions if matches_filter_func(tx)]

        all_transactions = []
        max_block_index = max((b.get('index', 0) for b in full_chain), default=0)

        for block in filtered_chain:
            for tx in block.get('transactions', []):
                tx_copy = tx.copy()
                tx_copy['block_index'] = block['index']
                tx_copy['block_timestamp'] = block['timestamp']
                tx_copy['_sort_key'] = block['index']
                all_transactions.append(tx_copy)

        for tx in filtered_pending:
            tx_copy = tx.copy()
            tx_copy['block_index'] = None
            tx_copy['block_timestamp'] = None
            tx_copy['_sort_key'] = max_block_index + 1
            all_transactions.append(tx_copy)

        sort = request.GET.get('sort', 'recent_to_oldest')
        if sort == 'recent_to_oldest':
            all_transactions.sort(key=lambda x: x['_sort_key'], reverse=True)
        elif sort == 'oldest_to_recent':
            all_transactions.sort(key=lambda x: x['_sort_key'])
        elif sort == 'highest_amount':
            all_transactions.sort(key=lambda x: x.get('amount', 0), reverse=True)
        elif sort == 'lowest_amount':
            all_transactions.sort(key=lambda x: x.get('amount', 0))

        total_transactions = len(all_transactions)

        paginator = Paginator(all_transactions, 10)
        page_number = request.GET.get('page', 1)
        page_obj = paginator.get_page(page_number)

        return render(request, 'blockchain.html', {
            'page_obj': page_obj,
            'total_blocks': len(filtered_chain),
            'total_transactions': total_transactions,
            'pending_transactions': filtered_pending,
        })
    except Exception as e:
        logger.error(f"Error fetching blockchain data: {str(e)}")
        messages.error(request, "Unable to retrieve blockchain data.")
        return redirect('donations')

@login_required
def download_ledger(request):
    full_chain = blockchain.get_chain()
    pending_transactions = blockchain.pending_transactions
    full_chain, pending_transactions = normalize_blockchain_data(full_chain, pending_transactions)

    search = request.GET.get('search', '').lower()
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    amount_min = request.GET.get('amount_min', '')
    amount_max = request.GET.get('amount_max', '')
    method = request.GET.get('method')

    matches_filter_func = get_matches_filter(search, date_from, date_to, amount_min, amount_max, method)

    filtered_chain = []
    for block in full_chain:
        filtered_txs = [tx for tx in block.get('transactions', []) if matches_filter_func(tx)]
        if filtered_txs:
            block_copy = block.copy()
            block_copy['transactions'] = filtered_txs
            filtered_chain.append(block_copy)

    filtered_pending = [tx for tx in pending_transactions if matches_filter_func(tx)]

    all_transactions = []
    max_block_index = max((b.get('index', 0) for b in full_chain), default=0)

    for block in filtered_chain:
        for tx in block.get('transactions', []):
            tx_copy = tx.copy()
            tx_copy['block_index'] = block['index']
            tx_copy['block_timestamp'] = block['timestamp']
            all_transactions.append(tx_copy)

    for tx in filtered_pending:
        tx_copy = tx.copy()
        tx_copy['block_index'] = None
        tx_copy['block_timestamp'] = None
        all_transactions.append(tx_copy)

    sort = request.GET.get('sort', 'recent_to_oldest')
    if sort == 'recent_to_oldest':
        all_transactions.sort(key=lambda x: x['block_index'] if x['block_index'] is not None else (max_block_index + 1), reverse=True)
    elif sort == 'oldest_to_recent':
        all_transactions.sort(key=lambda x: x['block_index'] if x['block_index'] is not None else (max_block_index + 1))
    elif sort == 'highest_amount':
        all_transactions.sort(key=lambda x: x.get('amount', 0), reverse=True)
    elif sort == 'lowest_amount':
        all_transactions.sort(key=lambda x: x.get('amount', 0))

    wb = Workbook()
    ws = wb.active
    ws.title = "Donation Ledger"

    headers = ['Block Index', 'Block Timestamp', 'Transaction ID', 'Donor', 'Email', 'Amount', 'Date', 'Method', 'Submitted By', 'Reviewed By']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for tx in all_transactions:
        donation_date = tx.get('donation_date')
        if isinstance(donation_date, str):
            try:
                donation_date = datetime.strptime(donation_date, '%Y-%m-%d').date()
            except:
                donation_date = None

        amount = float(tx.get('amount', 0)) if tx.get('amount') not in (None, '') else 0.0
        block_timestamp = tx.get('block_timestamp')
        if block_timestamp and block_timestamp.tzinfo is not None:
            block_timestamp = block_timestamp.replace(tzinfo=None)

        if donation_date and hasattr(donation_date, 'tzinfo') and donation_date.tzinfo is not None:
            donation_date = donation_date.replace(tzinfo=None)
        ws.append([
            tx.get('block_index', 'Pending') if tx.get('block_index') is not None else 'Pending',
            block_timestamp if block_timestamp else '',
            tx.get('transaction_id', ''),
            tx.get('donor', 'N/A'),
            tx.get('email', 'N/A'),
            amount,
            donation_date,
            tx.get('payment_method', 'N/A'),
            tx.get('submitted_by', 'N/A'),
            tx.get('reviewed_by', 'N/A')
        ])

    wrap_alignment = Alignment(wrap_text=True, vertical='top')

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = wrap_alignment

    for cell in ws['F'][1:]:
        cell.number_format = '"₱"#,##0.00'

    for cell in ws['G'][1:]:
        if cell.value:
            cell.number_format = 'YYYY-MM-DD'

    for cell in ws['B'][1:]:
        if cell.value:
            cell.number_format = 'YYYY-MM-DD HH:MM:SS'

    for column_cells in ws.columns:
        length = max(len(str(cell.value)) for cell in column_cells if cell.value)
        length = min(length + 4, 60)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = length

    ws.column_dimensions['D'].width = 35
    ws.column_dimensions['C'].width = 28
    ws.column_dimensions['E'].width = 32
    ws.column_dimensions['I'].width = 20
    ws.column_dimensions['J'].width = 20

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    response = HttpResponse(
        output,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="donation_ledger_{timestamp}.xlsx"'
    return response